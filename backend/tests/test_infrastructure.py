from __future__ import annotations

import asyncio
import threading
from collections.abc import Callable, Mapping
from io import BytesIO

import pytest
from conftest import ScriptedBackend, runtime
from docx import Document
from openpyxl import Workbook

from orion.chat.runtime import RequestCancelled
from orion.contracts import AssistantMessage, ModelToolCall, ModelTurn, RuntimeScope, ToolCall
from orion.integrations.infrastructure import LinuxFileMetadata, Target, TargetCatalog
from orion.tool_runtime.infrastructure import (
    _INFRASTRUCTURE_WORKERS,
    _blocking_handler,
    _linux_document_read,
    _linux_file_edit,
    infrastructure_definitions,
    infrastructure_registrations,
)
from orion.tool_runtime.registry import EXPAND_TOOL_NAME, ToolRegistryBuilder
from orion.tool_runtime.runner import ToolRunner


class FakeLinux:
    def __init__(self) -> None:
        self.calls = 0
        self.package = {"installed": False, "version": None}
        self.files: dict[str, bytes] = {"/etc/hosts": b"bounded text"}
        self.file_types: dict[str, str] = {}
        self.replacements: list[tuple[str, str]] = []
        self.writes = 0

    def inspect(
        self, target: Target, credential: object, sections: tuple[str, ...]
    ) -> Mapping[str, object]:
        self.calls += 1
        return {section: {"ok": True} for section in sections}

    def read_file(
        self, target: Target, credential: object, path: str, offset: int, length: int
    ) -> bytes:
        self.calls += 1
        return self.files.get(path, b"")[offset : offset + length]

    def file_metadata(self, target: Target, credential: object, path: str) -> LinuxFileMetadata:
        self.calls += 1
        return LinuxFileMetadata(
            self.file_types.get(path, "regular"), len(self.files.get(path, b""))
        )

    def write_file(self, target: Target, credential: object, path: str, content: bytes) -> None:
        self.calls += 1
        self.writes += 1
        self.files[path] = content

    def prepare_file_replacement(
        self, target: Target, credential: object, temporary_path: str, path: str
    ) -> None:
        self.calls += 1

    def replace_file(
        self, target: Target, credential: object, temporary_path: str, path: str
    ) -> None:
        self.calls += 1
        self.replacements.append((temporary_path, path))
        self.files[path] = self.files.pop(temporary_path)

    def remove_file(self, target: Target, credential: object, path: str) -> None:
        self.calls += 1
        self.files.pop(path, None)

    def service_status(
        self, target: Target, credential: object, service: str
    ) -> Mapping[str, object]:
        self.calls += 1
        return {"service": service, "load_state": "loaded", "active_state": "active"}

    def package_status(
        self, target: Target, credential: object, package: str
    ) -> Mapping[str, object]:
        self.calls += 1
        return self.package

    def service_preflight(self, target: Target, credential: object, service: str) -> None:
        self.calls += 1

    def restart_service(self, target: Target, credential: object, service: str) -> None:
        self.calls += 100

    def install_package(
        self, target: Target, credential: object, package: str, version: str | None
    ) -> None:
        self.calls += 100
        self.package = {"installed": True, "version": version or "1.0"}

    def health(self, target: Target, credential: object) -> None:
        pass


class BlockingPreflightLinux(FakeLinux):
    def __init__(self) -> None:
        super().__init__()
        self.reached = threading.Event()
        self.release = threading.Event()
        self.dispatches = 0

    def service_preflight(self, target: Target, credential: object, service: str) -> None:
        self.reached.set()
        assert self.release.wait(2)

    def restart_service(self, target: Target, credential: object, service: str) -> None:
        self.dispatches += 1


class BlockingVerificationLinux(FakeLinux):
    def __init__(self) -> None:
        super().__init__()
        self.verification_entered = threading.Event()
        self.release_verification = threading.Event()
        self.dispatches = 0

    def restart_service(self, target: Target, credential: object, service: str) -> None:
        self.dispatches += 1

    def service_status(
        self, target: Target, credential: object, service: str
    ) -> Mapping[str, object]:
        self.verification_entered.set()
        assert self.release_verification.wait(2)
        return {"service": service, "load_state": "loaded", "active_state": "active"}


class BoundedBlockingLinux(FakeLinux):
    def __init__(self) -> None:
        super().__init__()
        self.release = threading.Event()
        self.started = threading.Event()
        self._lock = threading.Lock()
        self.active = 0
        self.maximum_active = 0
        self.finished = 0

    def inspect(
        self, target: Target, credential: object, sections: tuple[str, ...]
    ) -> Mapping[str, object]:
        with self._lock:
            self.active += 1
            self.maximum_active = max(self.maximum_active, self.active)
            if self.active >= 8:
                self.started.set()
        assert self.release.wait(2)
        with self._lock:
            self.active -= 1
            self.finished += 1
        return {section: {"ok": True} for section in sections}

    def block(self) -> Mapping[str, object]:
        return self.inspect(_catalog().resolve("linux", "node"), "credential", ("cpu",))


class FakeGrafana:
    def dashboard_get(self, target: Target, credential: object, uid: str) -> Mapping[str, object]:
        return {"uid": uid, "url": "https://secret.example"}

    def datasource_query(
        self,
        target: Target,
        credential: object,
        uid: str,
        query: str,
        start: str,
        end: str,
        maximum: int,
    ) -> object:
        return {"series": [{"name": "cpu", "values": [1]}]}

    def alert_list(self, target: Target, credential: object) -> object:
        return [{"state": "firing", "name": "CPU"}]

    def annotation_create(
        self, target: Target, credential: object, payload: Mapping[str, object]
    ) -> str:
        return "42"

    def annotation_verify(
        self, target: Target, credential: object, annotation_id: str
    ) -> bool | None:
        return True

    def health(self, target: Target, credential: object) -> None:
        pass


class FakeZabbix:
    def __init__(self) -> None:
        self.ack_calls: list[dict[str, object]] = []
        self.acknowledged = False

    def call(
        self, target: Target, credential: object, method: str, params: Mapping[str, object]
    ) -> object:
        if method == "event.acknowledge":
            self.ack_calls.append(dict(params))
            self.acknowledged = True
            return {"eventids": params["eventids"]}
        if method == "event.get":
            return [
                {"eventid": event_id, "acknowledged": "1" if self.acknowledged else "0"}
                for event_id in params["eventids"]
            ]
        return [{"hostid": "1", "name": "host"}]

    def health(self, target: Target, credential: object) -> None:
        pass


def _catalog() -> TargetCatalog:
    return TargetCatalog.from_mapping(
        {
            "credentials": {"linux-key": "not-printed", "api": "not-printed"},
            "targets": {
                "linux": [
                    {
                        "target_ref": "node",
                        "display_name": "Node",
                        "credential_ref": "linux-key",
                        "host": "internal",
                    }
                ],
                "grafana": [
                    {
                        "target_ref": "observability",
                        "credential_ref": "api",
                        "base_url": "https://private",
                        "datasources": {"prom": "prometheus"},
                    }
                ],
                "zabbix": [
                    {
                        "target_ref": "monitoring",
                        "credential_ref": "api",
                        "base_url": "https://private",
                    }
                ],
            },
        }
    )


def _runner(linux: FakeLinux, grafana: FakeGrafana, zabbix: FakeZabbix) -> ToolRunner:
    builder = ToolRegistryBuilder()
    for registration in infrastructure_registrations(
        _catalog(), linux=linux, grafana=grafana, zabbix=zabbix
    ):
        builder.register(registration.definition, registration.handler)
    return ToolRunner(builder.freeze())


def _call(name: str, arguments: dict[str, object], call_id: str = "call") -> ModelToolCall:
    return ModelToolCall(call_id=call_id, tool_name=name, arguments=arguments)


def _scope() -> RuntimeScope:
    return RuntimeScope(session_id="session", principal_id="local", workspace_id="local")


async def _wait_for_thread_event(event: threading.Event) -> None:
    while not event.is_set():
        await asyncio.sleep(0.01)


async def _wait_for_task(task: asyncio.Task[object]) -> None:
    while not task.done():
        await asyncio.sleep(0.01)


async def _wait_for_count(predicate: Callable[[], bool]) -> None:
    for _ in range(200):
        if predicate():
            return
        await asyncio.sleep(0.01)
    raise AssertionError("timed out waiting for worker count")


def test_all_frozen_infrastructure_schemas_are_closed() -> None:
    names = {definition.name for definition in infrastructure_definitions()}
    assert names == {
        "linux.system.inspect",
        "linux.file.read",
        "linux.document.read",
        "linux.file.edit",
        "linux.service.status",
        "linux.package.status",
        "linux.service.restart",
        "linux.package.install",
        "grafana.dashboard.get",
        "grafana.datasource.query",
        "grafana.alert.list",
        "grafana.annotation.create",
        "zabbix.host.get",
        "zabbix.event.list",
        "zabbix.history.get",
        "zabbix.trigger.get",
        "zabbix.template.get",
        "zabbix.event.acknowledge",
    }
    assert all(
        definition.input_schema["additionalProperties"] is False
        for definition in infrastructure_definitions()
    )


@pytest.mark.anyio
async def test_unknown_target_is_rejected_before_credentials_or_executor() -> None:
    linux, grafana, zabbix = FakeLinux(), FakeGrafana(), FakeZabbix()
    result = await _runner(linux, grafana, zabbix).run_async(
        _call("linux.system.inspect", {"target_ref": "forged"}), _scope()
    )
    assert result.error is not None and result.error.code == "unknown_target"
    assert linux.calls == 0


@pytest.mark.anyio
async def test_successful_read_per_family_is_sanitized_and_source_bearing() -> None:
    linux, grafana, zabbix = FakeLinux(), FakeGrafana(), FakeZabbix()
    runner = _runner(linux, grafana, zabbix)
    results = [
        await runner.run_async(
            _call("linux.file.read", {"target_ref": "node", "path": "/etc/hosts"}), _scope()
        ),
        await runner.run_async(
            _call(
                "grafana.dashboard.get", {"target_ref": "observability", "dashboard_uid": "home"}
            ),
            _scope(),
        ),
        await runner.run_async(_call("zabbix.host.get", {"target_ref": "monitoring"}), _scope()),
    ]
    assert all(result.status == "success" and result.sources for result in results)
    assert {result.sources[0].source_kind for result in results} == {"linux", "grafana", "zabbix"}
    assert "private" not in str(results) and "not-printed" not in str(results)


def test_linux_document_tools_edit_text_through_verified_same_directory_temporary_file() -> None:
    linux = FakeLinux()
    linux.files["/tmp/orion-qa.txt"] = b"before\nafter\n"
    read = _linux_document_read(
        ToolCall(
            call_id="read",
            tool_name="linux.document.read",
            arguments={"target_ref": "node", "path": "/tmp/orion-qa.txt", "limit": 10},
            runtime_scope=_scope(),
        ),
        _catalog(),
        linux,
    )
    changed = _linux_file_edit(
        ToolCall(
            call_id="edit",
            tool_name="linux.file.edit",
            arguments={
                "target_ref": "node",
                "path": "/tmp/orion-qa.txt",
                "operations": [
                    {"kind": "replace_text", "old_text": "before", "new_text": "updated"}
                ],
            },
            runtime_scope=_scope(),
        ),
        _catalog(),
        linux,
    )

    assert read.status == "success" and read.data["format"] == "text"
    assert changed.status == "success" and changed.data["changed"] is True
    assert changed.data["verification"] == {"status": "verified", "format": "text", "operations": 1}
    assert linux.files["/tmp/orion-qa.txt"] == b"updated\nafter\n"
    assert linux.replacements and linux.replacements[0][0].startswith("/tmp/.orion-")


def test_linux_file_edit_rejects_ambiguous_or_cancelled_text_without_writing() -> None:
    linux = FakeLinux()
    linux.files["/tmp/orion-qa.txt"] = b"same same"
    arguments = {
        "target_ref": "node",
        "path": "/tmp/orion-qa.txt",
        "operations": [{"kind": "replace_text", "old_text": "same", "new_text": "new"}],
    }
    ambiguous = _linux_file_edit(
        ToolCall(
            call_id="ambiguous",
            tool_name="linux.file.edit",
            arguments=arguments,
            runtime_scope=_scope(),
        ),
        _catalog(),
        linux,
    )
    cancelled = _linux_file_edit(
        ToolCall(
            call_id="cancelled",
            tool_name="linux.file.edit",
            arguments={
                **arguments,
                "operations": [
                    {"kind": "replace_text", "old_text": "same same", "new_text": "new"}
                ],
            },
            runtime_scope=_scope(),
            cancellation_requested=lambda: True,
        ),
        _catalog(),
        linux,
    )

    assert ambiguous.error is not None and ambiguous.error.code == "ambiguous"
    assert cancelled.error is not None and cancelled.error.code == "cancelled"
    assert linux.files["/tmp/orion-qa.txt"] == b"same same"


def test_linux_file_edit_semantic_noop_does_not_write_or_replace() -> None:
    linux = FakeLinux()
    linux.files["/tmp/orion-qa.txt"] = b"unchanged"
    result = _linux_file_edit(
        ToolCall(
            call_id="noop",
            tool_name="linux.file.edit",
            arguments={
                "target_ref": "node",
                "path": "/tmp/orion-qa.txt",
                "operations": [
                    {"kind": "replace_text", "old_text": "unchanged", "new_text": "unchanged"}
                ],
            },
            runtime_scope=_scope(),
        ),
        _catalog(),
        linux,
    )

    assert result.status == "success" and result.data["changed"] is False
    assert linux.writes == 0 and not linux.replacements


@pytest.mark.parametrize("path", ["/tmp/orion-qa.docx", "/tmp/orion-qa.xlsx"])
def test_linux_office_semantic_noop_does_not_write_or_replace(path: str) -> None:
    payload = BytesIO()
    if path.endswith(".docx"):
        document = Document()
        document.add_paragraph("unchanged")
        document.save(payload)
        operation: dict[str, object] = {
            "kind": "set_paragraph",
            "paragraph_index": 0,
            "text": "unchanged",
        }
    else:
        workbook = Workbook()
        workbook.active.title = "Plan"
        workbook.active["A1"] = "unchanged"
        workbook.save(payload)
        operation = {"kind": "set_cell", "sheet": "Plan", "cell": "A1", "value": "unchanged"}
    linux = FakeLinux()
    linux.files[path] = payload.getvalue()
    result = _linux_file_edit(
        ToolCall(
            call_id="office-noop",
            tool_name="linux.file.edit",
            arguments={"target_ref": "node", "path": path, "operations": [operation]},
            runtime_scope=_scope(),
        ),
        _catalog(),
        linux,
    )

    assert result.status == "success" and result.data["changed"] is False
    assert linux.writes == 0 and not linux.replacements


@pytest.mark.parametrize("file_type", ["symlink", "directory", "other"])
def test_linux_file_edit_rejects_non_regular_files_before_mutation(file_type: str) -> None:
    linux = FakeLinux()
    linux.files["/tmp/orion-qa.txt"] = b"before"
    linux.file_types["/tmp/orion-qa.txt"] = file_type

    result = _linux_file_edit(
        ToolCall(
            call_id="nonregular",
            tool_name="linux.file.edit",
            arguments={
                "target_ref": "node",
                "path": "/tmp/orion-qa.txt",
                "operations": [{"kind": "replace_text", "old_text": "before", "new_text": "after"}],
            },
            runtime_scope=_scope(),
        ),
        _catalog(),
        linux,
    )

    assert result.error is not None and result.error.code == "unsupported_file_type"
    assert linux.files["/tmp/orion-qa.txt"] == b"before"
    assert not linux.replacements


def test_linux_file_edit_rejects_oversized_or_changed_final_file_without_mutation_claim() -> None:
    linux = FakeLinux()
    linux.files["/tmp/orion-qa.txt"] = b"before"
    linux.file_types["/tmp/orion-qa.txt"] = "regular"
    oversized = _linux_file_edit(
        ToolCall(
            call_id="oversized",
            tool_name="linux.file.edit",
            arguments={
                "target_ref": "node",
                "path": "/tmp/orion-qa.txt",
                "operations": [{"kind": "replace_text", "old_text": "before", "new_text": "after"}],
            },
            runtime_scope=_scope(),
        ),
        _catalog(),
        type(
            "OversizedLinux",
            (FakeLinux,),
            {
                "file_metadata": lambda self, *args: LinuxFileMetadata(
                    "regular", 4 * 1024 * 1024 + 1
                )
            },
        )(),
    )
    assert oversized.error is not None and oversized.error.code == "too_large"

    class ConcurrentChangeLinux(FakeLinux):
        def replace_file(self, target, credential, temporary_path, path):  # type: ignore[no-untyped-def]
            self.replacements.append((temporary_path, path))
            self.files[path] = b"after but not the validated bytes"
            self.files.pop(temporary_path, None)

    changed = ConcurrentChangeLinux()
    changed.files["/tmp/orion-qa.txt"] = b"before"
    result = _linux_file_edit(
        ToolCall(
            call_id="changed-final",
            tool_name="linux.file.edit",
            arguments={
                "target_ref": "node",
                "path": "/tmp/orion-qa.txt",
                "operations": [{"kind": "replace_text", "old_text": "before", "new_text": "after"}],
            },
            runtime_scope=_scope(),
        ),
        _catalog(),
        changed,
    )
    assert result.error is not None and result.error.code == "outcome_unknown"


@pytest.mark.anyio
async def test_cancellation_between_preflight_and_dispatch_issues_no_restart() -> None:
    linux, grafana, zabbix = FakeLinux(), FakeGrafana(), FakeZabbix()
    cancelled = False

    def cancel_after_preflight() -> bool:
        return cancelled or linux.calls >= 1

    result = await _runner(linux, grafana, zabbix).run_async(
        _call("linux.service.restart", {"target_ref": "node", "service": "nginx"}),
        _scope(),
        cancel_after_preflight,
    )
    assert result.error is not None and result.error.code == "cancelled"
    assert linux.calls == 1


@pytest.mark.anyio
async def test_package_convergence_and_zabbix_acknowledgement_dispatch_once() -> None:
    linux, grafana, zabbix = FakeLinux(), FakeGrafana(), FakeZabbix()
    runner = _runner(linux, grafana, zabbix)
    install = await runner.run_async(
        _call("linux.package.install", {"target_ref": "node", "package": "curl"}), _scope()
    )
    ack = await runner.run_async(
        _call("zabbix.event.acknowledge", {"target_ref": "monitoring", "event_ids": ["1", "2"]}),
        _scope(),
    )
    assert install.status == "success" and install.data["changed"] is True
    assert ack.status == "success" and ack.data["changed"] is True
    assert len(zabbix.ack_calls) == 1
    assert zabbix.ack_calls[0]["eventids"] == ["1", "2"]


@pytest.mark.anyio
async def test_runtime_cancellation_is_observed_between_linux_preflight_and_dispatch(store) -> None:  # type: ignore[no-untyped-def]
    linux, grafana, zabbix = BlockingPreflightLinux(), FakeGrafana(), FakeZabbix()
    builder = ToolRegistryBuilder()
    for registration in infrastructure_registrations(
        _catalog(), linux=linux, grafana=grafana, zabbix=zabbix
    ):
        builder.register(registration.definition, registration.handler)
    backend = ScriptedBackend(
        [
            ModelTurn(
                tool_calls=(
                    ModelToolCall(
                        call_id="expand",
                        tool_name=EXPAND_TOOL_NAME,
                        arguments={"tool_names": ["linux.service.restart"]},
                    ),
                )
            ),
            ModelTurn(
                tool_calls=(
                    ModelToolCall(
                        call_id="restart",
                        tool_name="linux.service.restart",
                        arguments={"target_ref": "node", "service": "nginx"},
                    ),
                )
            ),
            ModelTurn(assistant=AssistantMessage(content="fallback")),
        ]
    )
    chat = runtime(store, backend, builder.freeze())
    session_id = store.create_session()
    request_id = chat.begin(session_id, "restart")
    task = asyncio.create_task(chat.run(session_id, request_id))
    await asyncio.wait_for(_wait_for_thread_event(linux.reached), 2)
    assert chat.cancel(request_id)
    linux.release.set()
    await asyncio.wait_for(_wait_for_task(task), 2)
    with pytest.raises(RequestCancelled):
        task.result()
    assert linux.dispatches == 0
    assert store.request(request_id)["status"] == "cancelled"
    tool_result = next(
        item.payload["result"]
        for item in store.timeline(session_id)
        if item.kind == "tool_result" and item.tool_name == "linux.service.restart"
    )
    assert tool_result["error"]["code"] == "cancelled"


@pytest.mark.anyio
async def test_runtime_cancellation_after_restart_preserves_verified_dispatch_result(store) -> None:  # type: ignore[no-untyped-def]
    linux, grafana, zabbix = BlockingVerificationLinux(), FakeGrafana(), FakeZabbix()
    builder = ToolRegistryBuilder()
    for registration in infrastructure_registrations(
        _catalog(), linux=linux, grafana=grafana, zabbix=zabbix
    ):
        builder.register(registration.definition, registration.handler)
    backend = ScriptedBackend(
        [
            ModelTurn(
                tool_calls=(
                    ModelToolCall(
                        call_id="expand",
                        tool_name=EXPAND_TOOL_NAME,
                        arguments={"tool_names": ["linux.service.restart"]},
                    ),
                )
            ),
            ModelTurn(
                tool_calls=(
                    ModelToolCall(
                        call_id="restart",
                        tool_name="linux.service.restart",
                        arguments={"target_ref": "node", "service": "nginx"},
                    ),
                )
            ),
            ModelTurn(assistant=AssistantMessage(content="must not be replayed")),
        ]
    )
    chat = runtime(store, backend, builder.freeze())
    session_id = store.create_session()
    request_id = chat.begin(session_id, "restart")
    task = asyncio.create_task(chat.run(session_id, request_id))

    await asyncio.wait_for(_wait_for_thread_event(linux.verification_entered), 2)
    assert linux.dispatches == 1
    assert chat.cancel(request_id)
    linux.release_verification.set()

    await asyncio.wait_for(_wait_for_task(task), 2)
    with pytest.raises(RequestCancelled):
        task.result()

    assert linux.dispatches == 1
    assert len(backend.calls) == 2
    assert store.request(request_id)["status"] == "cancelled"
    tool_result = next(
        item.payload["result"]
        for item in store.timeline(session_id)
        if item.kind == "tool_result" and item.tool_name == "linux.service.restart"
    )
    assert tool_result["status"] == "success"
    assert tool_result["data"]["changed"] is True
    assert tool_result["data"]["verification"]["status"] == "verified"
    assert "rollback" not in str(tool_result).lower()


@pytest.mark.anyio
async def test_infrastructure_workers_are_bounded_and_capacity_is_restored() -> None:
    linux = BoundedBlockingLinux()
    handler = _blocking_handler(lambda call: linux.block())
    for index in range(9):
        asyncio.create_task(
            handler(
                ToolCall(
                    call_id=str(index),
                    tool_name="linux.system.inspect",
                    arguments={},
                    runtime_scope=_scope(),
                )
            )
        )
    await asyncio.wait_for(_wait_for_thread_event(linux.started), 2)
    await asyncio.sleep(0.05)
    assert linux.maximum_active == 8
    assert linux.finished == 0
    linux.release.set()
    await asyncio.wait_for(_wait_for_count(lambda: linux.finished == 9), 2)
    assert linux.finished == 9
    workers = _INFRASTRUCTURE_WORKERS[asyncio.get_running_loop()]
    await _wait_for_count(lambda: workers._value == 8)
    assert workers._value == 8  # noqa: SLF001 - capacity restoration invariant.
    follow_up = asyncio.create_task(
        handler(
            ToolCall(
                call_id="after",
                tool_name="linux.system.inspect",
                arguments={},
                runtime_scope=_scope(),
            )
        )
    )
    await asyncio.wait_for(_wait_for_count(lambda: linux.finished == 10), 2)
    await _wait_for_count(lambda: workers._value == 8)
    assert workers._value == 8
    assert follow_up.done()
    assert linux.finished == 10
